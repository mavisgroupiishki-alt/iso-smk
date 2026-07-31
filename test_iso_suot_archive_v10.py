"""Regression tests for ISO/SUOT archive extraction v10.

Run from the project root:
    python test_iso_suot_archive_v10.py
"""
from __future__ import annotations

import importlib.util
import io
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

spec = importlib.util.spec_from_file_location("iso_server_v10", ROOT / "server.py")
server = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(server)


def make_docx(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сотрудники"
    sheet.append(["ФИО", "Должность", "Удостоверение ОТ"])
    sheet.append(["Иванов Иван Иванович", "Прораб", "№123 от 01.02.2026"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_archive() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "исо тест/Перечень поставщиков.docx",
            make_docx("Перечень поставщиков: Кабель — ООО Альфа"),
        )
        archive.writestr(
            "исо тест/Перечень объектов.docx",
            make_docx("Перечень объектов: ремонт здания по адресу г. Минск"),
        )
        archive.writestr("исо тест/Список сотрудников.xlsx", make_xlsx())
        archive.writestr(
            "исо тест/Реквизиты.docx",
            make_docx(
                "Заказчик: ООО Тест. УНП 123456789. Юридический адрес: г. Минск. "
                "Банковские реквизиты: р/с BY00TEST, БИК TESTBY2X"
            ),
        )
        archive.writestr("исо тест/Удостоверение по охране труда Иванов.pdf", b"fake pdf")
    return buffer.getvalue()


server._simple_ai_call = lambda *args, **kwargs: (
    "Название: Тест\n"
    "УНП: 123456789\n"
    "Юридический адрес: г. Минск\n"
    "Банковские реквизиты: р/с BY00TEST, БИК TESTBY2X\n"
    "Директор: Иванов И.И.\n"
    "Телефон/email: не найдено"
)
server.vision_extract = lambda data, filename, api_key, *args, **kwargs: (
    "Удостоверение по охране труда: Иванов Иван Иванович, №123 от 01.02.2026"
)

result = server.extract_archive_with_vision(build_archive(), "iso-test.zip", "test-key")
text = result["text"] if isinstance(result, dict) else result

required = [
    "ИСХОДНЫЕ ДОКУМЕНТЫ ИЗ АРХИВА",
    "ООО Альфа",
    "ремонт здания",
    "Иванов Иван Иванович",
    "Удостоверение по охране труда",
    "Реквизиты компании",
]
missing = [item for item in required if item not in text]
assert not missing, f"Archive data was lost: {missing}\n\n{text}"
print("OK: suppliers, objects, staff, OT certificate and requisites are all preserved")
