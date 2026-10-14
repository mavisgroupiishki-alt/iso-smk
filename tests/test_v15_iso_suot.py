import io
import re
import sys
import time
import unicodedata
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import generator
import server
from generator_iso_suot_templates import generate_iso_suot_package_v2
from docx import Document
from docx.enum.section import WD_ORIENT


def xml_text(doc_bytes):
    with zipfile.ZipFile(io.BytesIO(doc_bytes)) as z:
        return '\n'.join(
            z.read(n).decode('utf-8', 'ignore')
            for n in z.namelist()
            if n.startswith('word/') and n.endswith('.xml')
        )


def sample_data():
    return {
        'company': {
            'name': 'ОмиТрейд', 'form': 'ООО', 'unp': '790798634',
            'address': '212001, г. Могилев, ул. Кирова, 26Б/4', 'city': 'Могилев',
            'director_fio': 'Глушинский Олег Иванович', 'director_position': 'Директор',
            'scope': 'СТАРАЯ ОБЛАСТЬ КОМПАНИИ',
        },
        'certification': {
            'standard': 'iso_suot',
            'scope': 'Текущий ремонт и отделочные работы',
            'audit_date': '20.08.2026',
        },
        'staff': [
            {'fio':'Глушинский Олег Иванович','position':'Директор','is_worker':False,'ot_certificate':True},
            {'fio':'Лукашик Сергей Васильевич','position':'Заместитель директора-главный инженер','is_worker':False,'ot_certificate':True,'employment_type':'совместительство'},
            {'fio':'Волков Василий Анатольевич','position':'Производитель работ','is_worker':False,'ot_certificate':True},
            {'fio':'Грук Сергей Борисович','position':'Штукатур','is_worker':True},
        ],
        'workers': ['Штукатур'],
        'suppliers': [{'name':'РБК','type':'Плиты, профиль, смеси'}],
        'objects': [{'name':'Текущий ремонт поликлиники','year':'2026'}],
    }


def test_fast_deterministic_iso_suot_package():
    data = sample_data()
    company = dict(data['company'])
    company['scope'] = data['certification']['scope']
    itr = [x for x in data['staff'] if not x['is_worker']]
    workers = [x for x in data['staff'] if x['is_worker']]
    dates = generator.calculate_dates(data['certification']['audit_date'])
    resp = generator.select_responsible(itr)
    started = time.time()
    result = generate_iso_suot_package_v2(
        company, itr, dates, resp, product='iso_suot', workers=workers,
        suppliers=data['suppliers'], objects=data['objects']
    )
    assert time.time() - started < 10
    assert result['docs']
    assert not result.get('warnings')
    assert not any('Варта' in d['name'] for d in result['docs'])
    combined = '\n'.join(xml_text(d['bytes']) for d in result['docs'])
    assert 'Варта' not in combined
    assert '[Наименование]' not in combined
    assert '[Форма_собственности_сокращенная]' not in combined
    assert 'СТАРАЯ ОБЛАСТЬ КОМПАНИИ' not in combined
    assert 'Текущий ремонт и отделочные работы' in combined
    assert 'Лукашик Сергей Васильевич' in combined
    assert 'Волков Василий Анатольевич' in combined
    assert any('ИОТ Штукатур' in d['name'] for d in result['docs'])
    assert any('Карточка оценки поставщика' in d['name'] for d in result['docs'])
    paths = {doc['archive_path'] for doc in result['docs']}
    assert any(path.startswith('ИСО 9001/1 Приказы/') for path in paths)
    assert any(path.startswith('СУОТ/1 Приказы/') for path in paths)
    assert any(path.startswith('СУОТ/5 Риски/') for path in paths)
    assert all('/' in path and '..' not in path for path in paths)
    org_structure = next(d for d in result['docs'] if d['name'] == 'ОмиТрейд - Приложение 1 организационная структура.docx')
    document_list = next(d for d in result['docs'] if d['name'] == 'ОмиТрейд - Приложение 2 Перечень документов СМК.docx')
    assert 'Лукашик Сергей Васильевич' in xml_text(org_structure['bytes'])
    assert 'Грук Сергей Борисович' in xml_text(org_structure['bytes'])
    assert 'СТП СМК Управление рисками' in xml_text(document_list['bytes'])
    assert not any('валидац' in d['name'].lower() for d in result['docs'])


def test_iso_suot_package_includes_professional_diseases_journal():
    """The statutory professional-diseases journal must stay in initial SUOT output."""
    data = sample_data()
    result = generator.generate_package(data, 'dummy', 'iso_suot')

    journal = next(
        doc for doc in result['docs']
        if 'журнал регистрации профзаболеван' in doc['name'].lower()
    )
    rendered = Document(io.BytesIO(journal['bytes']))
    text = unicodedata.normalize(
        'NFC',
        ' '.join('\n'.join(paragraph.text for paragraph in rendered.paragraphs).split()),
    )
    assert 'Журнал регистрации профессиональных заболеваний' in text
    assert 'ОмиТрейд' in text
    assert 'Варта' not in text


def test_suot_report_has_current_scope_and_current_staff_count():
    data = sample_data()
    company = dict(data['company']); company['scope'] = data['certification']['scope']
    itr = [x for x in data['staff'] if not x['is_worker']]
    workers = [x for x in data['staff'] if x['is_worker']]
    result = generate_iso_suot_package_v2(
        company, itr, generator.calculate_dates('20.08.2026'), generator.select_responsible(itr),
        product='suot', workers=workers
    )
    report = next(d for d in result['docs'] if 'Отчет для анализа OH&S' in d['name'])
    source = next(d for d in result['docs'] if 'Исходная информация СУОТ' in d['name'])
    assert data['certification']['scope'] in xml_text(report['bytes'])
    source_xml = xml_text(source['bytes'])
    assert '4 чел.' in source_xml or '>4 чел.<' in source_xml
    assert '17 человек' not in source_xml


def test_periodika_iso_contains_only_annual_documents_and_current_objects():
    data = sample_data()
    company = dict(data['company']); company['scope'] = data['certification']['scope']
    itr = [x for x in data['staff'] if not x['is_worker']]
    workers = [x for x in data['staff'] if x['is_worker']]
    result = generate_iso_suot_package_v2(
        company, itr, generator.calculate_dates('20.08.2026'), generator.select_responsible(itr),
        product='iso', periodika=True, workers=workers,
        suppliers=data['suppliers'], objects=data['objects']
    )
    names = [doc['name'] for doc in result['docs']]
    assert any('Политика в области качества' in name for name in names)
    assert any('РЕЕСТР РИСКОВ' in name for name in names)
    assert any('оценке удовлетворенности заказчиков' in name.lower() for name in names)
    assert any('сводн' in name.lower() for name in names)
    assert not any('РК СМК' in name or 'СТП СМК' in name for name in names)
    assert not any('Приказ' in name or 'ДИ ' in name for name in names)
    assert not any('Журнал' in name for name in names)
    assert not any('Приложение 1 организационная структура' in name for name in names)
    assert not any('Приложение 2 Перечень документов СМК' in name for name in names)
    report = next(doc for doc in result['docs'] if 'оценке удовлетворенности заказчиков' in doc['name'].lower())
    assert 'Текущий ремонт поликлиники' in xml_text(report['bytes'])


def test_periodika_suot_excludes_baseline_manuals_instructions_and_orders():
    data = sample_data()
    company = dict(data['company']); company['scope'] = data['certification']['scope']
    itr = [x for x in data['staff'] if not x['is_worker']]
    workers = [x for x in data['staff'] if x['is_worker']]
    result = generate_iso_suot_package_v2(
        company, itr, generator.calculate_dates('20.08.2026'), generator.select_responsible(itr),
        product='suot', periodika=True, workers=workers
    )
    names = [doc['name'] for doc in result['docs']]
    assert any('цели-и-мероприятия' in name for name in names)
    assert any('SWOT анализ OH&S' in name for name in names)
    assert any('Отчет для анализа OH&S' in name for name in names)
    assert any('Программа внутренних аудитов' in name for name in names)
    assert not any('Руководство СУОТ' in name or 'СТП 8.1' in name for name in names)
    assert not any('ИНСТРУКЦИЯ' in name or 'ИОТ ' in name for name in names)
    assert not any('Приказ' in name or 'Журнал' in name for name in names)


def test_generator_uses_periodika_mode_from_card_data():
    data = sample_data()
    data['certification']['package_mode'] = 'periodika'
    result = generator.generate_package(data, 'dummy', 'iso')
    names = [doc['name'] for doc in result['docs']]
    assert names
    assert any('Политика в области качества' in name for name in names)
    assert not any('РК СМК' in name or 'СТП СМК' in name for name in names)
    assert not any('Приказ' in name for name in names)


def test_generator_refuses_iso_package_with_missing_critical_card_fields():
    data = sample_data()
    data['certification'].pop('scope')
    data['company'].pop('scope')
    data['dates'] = {}
    data['certification'].pop('audit_date')
    data['staff'] = []

    result = generator.generate_package(data, 'dummy', 'iso')

    assert result['docs'] == []
    assert 'область применения' in result['error']
    assert 'дата оформления документов' in result['error']
    assert 'актуальный штат' in result['error']


def test_generator_does_not_treat_template_staff_row_as_a_real_employee():
    data = sample_data()
    data['staff'] = [{'fio': 'Не указано Должность', 'position': 'ИТР'}]

    result = generator.generate_package(data, 'dummy', 'iso')

    assert result['docs'] == []
    assert 'актуальный штат' in result['error']


def test_initial_suot_keeps_orders_and_instructions_without_periodika_flag():
    data = sample_data()
    result = generator.generate_package(data, 'dummy', 'suot')
    assert not result.get('error')
    names = [doc['name'] for doc in result['docs']]
    assert not result.get('error')

    assert any('Приказ 4-OH&S' in name for name in names)
    assert any('Перечень инструкций по ОТ' in name for name in names)
    assert any('ИОТ Штукатур' in name for name in names)


def test_initial_suot_has_the_complete_staff_dependent_control_set():
    """The fast renderer must not drop documents that were excluded as stale templates.

    The approved Bayruz package contains these documents.  They must be rebuilt
    from the current card, rather than copied with former employees or omitted.
    """
    data = sample_data()
    result = generator.generate_package(data, 'dummy', 'suot')
    names = [doc['name'] for doc in result['docs']]

    required = [
        'Реестр неприемлемых рисков',
        'Карта рисков Глушинский Олег Иванович',
        'Карта рисков Лукашик Сергей Васильевич',
        'Карта рисков Волков Василий Анатольевич',
        'Карта рисков Грук Сергей Борисович',
        'Протокол проверки знаний',
        'Контрольный лист внутреннего аудита',
        'Отчет по внутреннему аудиту',
        'Карточка учета выдачи СИЗ Грук Сергей Борисович',
        'Акт ежеквартального контроля за 2 квартал 2026',
    ]
    for expected in required:
        assert any(expected in name for name in names), expected
    assert sum('Программа управления' in name for name in names) == 1
    assert not any('12. реестр рисков и мероприятия' in name for name in names)


def test_suot_appointment_orders_have_the_required_operational_clauses():
    """Orders 4–7 must retain the duties from the approved Bayruz reference."""
    data = sample_data()
    result = generator.generate_package(data, 'dummy', 'suot')
    def visible(doc_bytes):
        document = Document(io.BytesIO(doc_bytes))
        return '\n'.join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
    docs = {doc['name']: visible(doc['bytes']) for doc in result['docs']}

    order4 = next(text for name, text in docs.items() if 'Приказ 4-OH&S' in name)
    order5 = next(text for name, text in docs.items() if 'Приказ 5-OH&S' in name)
    order6 = next(text for name, text in docs.items() if 'Приказ 6-OH&S' in name)
    order7 = next(text for name, text in docs.items() if 'Приказ 7-OH&S' in name)

    assert 'вводного инструктажа' in order4
    assert 'электрохозяйства' in order4
    assert 'комиссии для проверки знаний' in order4
    assert 'периодических испытаний и проверок ручных электрических машин' in order4
    assert 'пожарно-технической комиссии' in order4
    assert 'общих технических осмотров зданий и сооружений' in order4
    assert 'председатель — Директор — Глушинский Олег Иванович; члены комиссии — Заместитель директора-главный инженер — Лукашик Сергей Васильевич; Производитель работ — Волков Василий Анатольевич.' in order4
    assert 'члены комиссии — Директор — Глушинский Олег Иванович' not in order4
    assert 'представление требуемой информации' in order5
    assert 'необходимые условия для проведения аудита' in order5
    assert 'С целью проведения внутреннего аудита' in order5
    assert 'реестр неприемлемых рисков' in order6
    assert 'мероприятия по снижению профессиональных рисков' in order6
    assert 'Для проведения работ по актуализации карт' in order6
    assert 'журнале учета выдачи инструкций' in order7
    assert 'Общее руководство разработкой инструкций' in order7
    assert 'В целях обеспечения учёта, актуализации и выдачи инструкций' in order7


def test_suot_training_and_safety_day_records_use_only_current_staff():
    """Do not retain Varta's duplicate "главный специалист" training records."""
    data = sample_data()
    result = generator.generate_package(data, 'dummy', 'suot')
    docs = {doc['name']: _visible_doc_text(doc['bytes']) for doc in result['docs']}

    training_order = next(text for name, text in docs.items() if 'Приказ 10.1-OH&S' in name)
    training_protocol = next(text for name, text in docs.items() if 'Протокол внутреннего обучения' in name)
    training_program = next(text for name, text in docs.items() if 'Программа внутреннего обучения' in name)
    safety_day = next(text for name, text in docs.items() if 'Приказ 9-OH&S' in name)

    for text in (training_order, training_protocol, training_program, safety_day):
        assert 'Главный специалист' not in text
        assert 'Лукашик Сергей Васильевич' in text
        assert 'Волков Василий Анатольевич' in text
    assert 'председатель — Директор — Глушинский Олег Иванович; члены комиссии — Заместитель директора-главный инженер — Лукашик Сергей Васильевич; Производитель работ — Волков Василий Анатольевич.' in safety_day


def test_suot_never_adds_worker_instructions_absent_from_the_staff_roster():
    """Bayruz regression: the source roster has no construction workers."""
    data = sample_data()
    data['staff'] = [
        {'fio': 'Яскин Илья Васильевич', 'position': 'Директор', 'is_worker': False, 'ot_certificate': True},
        {'fio': 'Падуто Александр Сергеевич', 'position': 'Начальник участка', 'is_worker': False, 'ot_certificate': True},
        {'fio': 'Кульбакин Александр Валерьевич', 'position': 'Бухгалтер', 'is_worker': False, 'employment_type': 'совместительство'},
    ]
    data['workers'] = []
    result = generator.generate_package(data, 'dummy', 'suot')
    names = '\n'.join(doc['name'] for doc in result['docs']).lower()

    assert 'бетонщик' not in names
    assert 'каменщик' not in names
    assert 'монтажник строительных конструкций' not in names
    assert 'карточка учета выдачи сиз (бланк)' in names
    assert 'карточка учета выдачи сиз яскин' not in names
    assert 'карточка учета выдачи сиз падуто' not in names


def test_suot_commission_prefers_the_actual_operational_manager_to_accountant():
    """Bayruz regression: a two-person commission must not promote the accountant."""
    itr = [
        {'fio': 'Яскин Илья Васильевич', 'position': 'Директор'},
        {'fio': 'Падуто Александр Сергеевич', 'position': 'Начальник участка'},
        {'fio': 'Кульбакин Александр Валерьевич', 'position': 'Бухгалтер'},
    ]
    responsible = generator.select_responsible(itr)
    assert [person['fio'] for person in responsible['auditors']] == [
        'Яскин Илья Васильевич',
        'Падуто Александр Сергеевич',
    ]


def test_suot_audit_checklist_never_assigns_a_person_to_audit_themself():
    """Bayruz regression: director/auditor names must not be the same."""
    data = sample_data()
    data['company']['director_fio'] = 'Яскин Илья Васильевич'
    data['company']['director_position'] = 'Директор'
    data['staff'] = [
        {'fio': 'Яскин Илья Васильевич', 'position': 'Директор', 'is_worker': False, 'ot_certificate': True},
        {'fio': 'Падуто Александр Сергеевич', 'position': 'Начальник участка', 'is_worker': False, 'ot_certificate': True},
        {'fio': 'Кульбакин Александр Валерьевич', 'position': 'Бухгалтер', 'is_worker': False},
    ]
    result = generator.generate_package(data, 'dummy', 'suot')
    checklist = next(doc['bytes'] for doc in result['docs']
                     if 'Контрольный лист внутреннего аудита 1' in doc['name'])
    document = Document(io.BytesIO(checklist))
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)

    assert 'Проверяемое лицо: Директор — Яскин Илья Васильевич' in text
    assert 'Аудитор: Начальник участка — Падуто Александр Сергеевич' in text


def test_suot_audit_program_uses_another_itr_when_only_director_has_ot_certificate():
    """A valid second ITR is safer than a placeholder in a finished programme."""
    data = sample_data()
    data['company']['director_fio'] = 'Яскин Илья Васильевич'
    data['company']['director_position'] = 'Директор'
    data['staff'] = [
        {'fio': 'Яскин Илья Васильевич', 'position': 'Директор', 'is_worker': False, 'ot_certificate': True},
        {'fio': 'Падуто Александр Сергеевич', 'position': 'Начальник участка', 'is_worker': False},
        {'fio': 'Кульбакин Александр Валерьевич', 'position': 'Бухгалтер', 'is_worker': False},
    ]

    result = generator.generate_package(data, 'dummy', 'suot')
    programme = next(doc['bytes'] for doc in result['docs'] if 'Программа внутренних аудитов' in doc['name'])
    rows = Document(io.BytesIO(programme)).tables[0].rows
    director_row = next(row for row in rows if row.cells[1].text.startswith('Директор'))

    assert 'Падуто Александр Сергеевич' in director_row.cells[4].text
    assert 'ТРЕБУЕТ УТОЧНЕНИЯ' not in director_row.cells[4].text


def test_suot_role_orders_never_assign_operational_safety_to_an_accountant():
    """Operational duties follow the actual operational manager, not office fallback roles."""
    data = sample_data()
    data['company']['director_fio'] = 'Яскин Илья Васильевич'
    data['company']['director_position'] = 'Директор'
    data['staff'] = [
        {'fio': 'Яскин Илья Васильевич', 'position': 'Директор', 'is_worker': False, 'ot_certificate': True},
        {'fio': 'Падуто Александр Сергеевич', 'position': 'Начальник участка', 'is_worker': False},
        {'fio': 'Кульбакин Александр Валерьевич', 'position': 'Бухгалтер', 'is_worker': False},
    ]

    result = generator.generate_package(data, 'dummy', 'suot')
    docs = {doc['name']: _visible_doc_text(doc['bytes']) for doc in result['docs']}
    order4 = next(text for name, text in docs.items() if 'Приказ 4-OH&S' in name)
    order7 = next(text for name, text in docs.items() if 'Приказ 7-OH&S' in name)

    for text in (order4, order7):
        assert 'Падуто Александр Сергеевич' in text
        assert 'Кульбакин Александр Валерьевич' not in text


def test_suot_risk_program_uses_only_current_worker_professions():
    """Old template professions must never leak into a new company's risk programme."""
    data = sample_data()
    data['workers'] = ['Монтажник связи', 'Электромонтажник', 'Подсобный рабочий']
    data['staff'] = [person for person in data['staff'] if not person['is_worker']]

    result = generator.generate_package(data, 'dummy', 'suot')
    programme = next(doc['bytes'] for doc in result['docs'] if 'Программа управления охраной труда' in doc['name'])
    text = _visible_doc_text(programme)

    for profession in data['workers']:
        assert profession in text
    for stale in ('Штукатур', 'Кровельщик', 'Электрогазосварщик'):
        assert stale not in text


def test_suot_lists_and_risk_cards_use_manual_worker_professions_without_template_roles():
    """Manual professions are real package data, not placeholders or sample-template rows."""
    data = sample_data()
    workers = ['Монтажник стальных конструкций', 'Электромонтажник', 'Разнорабочий']
    data['workers'] = workers
    data['staff'] = [person for person in data['staff'] if not person['is_worker']]

    result = generator.generate_package(data, 'dummy', 'suot')
    docs = {doc['name']: doc['bytes'] for doc in result['docs']}
    current_list = _visible_doc_text(next(value for name, value in docs.items() if '6 Перечень рабочих мест' in name))
    ppe_list = _visible_doc_text(next(value for name, value in docs.items() if '12 Перечень СИЗ' in name))
    current_staff_lists = [
        _visible_doc_text(value)
        for name, value in docs.items()
        if any(f'СУОТ - {number} ' in name for number in (4, 5, 6, 7, 8, 9, 11, 12))
    ]

    for profession in workers:
        assert profession in current_list
        assert profession in ppe_list
        assert any(f'Карта рисков для профессии {profession}' in name for name in docs)
    for stale_role in ('Техник по связи', 'Инженер-проектировщик', 'Электромонтер ОПС'):
        assert stale_role not in current_list
        assert stale_role not in ppe_list
        assert all(stale_role not in text for text in current_staff_lists)


def test_suot_manual_worker_professions_receive_ppe_cards_and_pass_manifest():
    """Selected professions have no FIO yet, but still require a PPE card."""
    data = sample_data()
    workers = ['Монтажник стальных конструкций', 'Электромонтажник', 'Разнорабочий']
    data['workers'] = workers
    data['staff'] = [person for person in data['staff'] if not person['is_worker']]

    result = generator.generate_package(data, 'dummy', 'suot')
    docs = {doc['name']: doc['bytes'] for doc in result['docs']}

    assert not result.get('error')
    for profession in workers:
        name = next(name for name in docs if f'Карточка учета выдачи СИЗ {profession}' in name)
        text = _visible_doc_text(docs[name])
        assert f'Профессия: {profession}' in text
        assert 'ФИО: ТРЕБУЕТ УТОЧНЕНИЯ' not in text


def test_suot_archive_routes_knowledge_forms_and_electrical_instruction_to_working_folders():
    data = sample_data()
    result = generator.generate_package(data, 'dummy', 'suot')
    paths = {doc['name']: doc['archive_path'] for doc in result['docs']}

    protocol_path = next(path for name, path in paths.items() if 'Протокол проверки знаний' in name)
    electrical_instruction_path = next(path for name, path in paths.items() if '7 при работе с электроинструментом' in name)

    assert protocol_path.startswith('СУОТ/6 Проверка знаний/')
    assert electrical_instruction_path.startswith('СУОТ/3 Инструкции/')


def test_suot_archive_route_normalizes_decomposed_cyrillic_filenames():
    """Archive folders must not depend on how macOS encodes the letter «й»."""
    from generator_iso_suot_templates import _archive_folder_for_document

    filename = 'БАЙРУЗ СУОТ - ВОПРОСЫ по проверке знаний.docx'
    assert _archive_folder_for_document(filename) == 'СУОТ/6 Проверка знаний'


def test_suot_template_lists_replace_generic_specialist_and_current_document_year():
    data = sample_data()
    result = generator.generate_package(data, 'dummy', 'suot')
    docs = {doc['name']: _visible_doc_text(doc['bytes']) for doc in result['docs']}

    combined = '\n'.join(docs.values())
    assert 'Главный специалист' not in combined
    assert 'Заместитель директора-главный инженер' in combined

    stp = next(value for name, value in docs.items() if 'СТП 8.1-02' in name)
    assert 'СТП OH&S 8.1-02- 2025' not in stp
    assert 'СТП OH&S 8.1-02-2026' in stp


def test_suot_audit_program_is_landscape_and_uses_selected_document_date():
    data = sample_data()
    data['dates'] = {'document_date': '14.10.2026'}
    result = generator.generate_package(data, 'dummy', 'suot')

    audit = next(doc for doc in result['docs'] if 'Программа внутренних аудитов' in doc['name'])
    order = next(doc for doc in result['docs'] if 'Приказ 4-OH&S' in doc['name'])
    audit_document = Document(io.BytesIO(audit['bytes']))

    assert audit_document.sections[0].orientation == WD_ORIENT.LANDSCAPE
    assert '14.10.2026' in _visible_doc_text(audit['bytes'])
    assert '14.10.2026' in _visible_doc_text(order['bytes'])


def test_generator_uses_only_company_director_in_suot_orders():
    data = sample_data()
    data['staff'].append({
        'fio': 'Ошибочный Директор Иванович', 'position': 'Директор',
        'is_worker': False, 'ot_certificate': True,
    })

    result = generator.generate_package(data, 'dummy', 'suot')
    combined = '\n'.join(xml_text(doc['bytes']) for doc in result['docs'] if 'Приказ' in doc['name'])

    assert 'Глушинский Олег Иванович' in combined
    assert 'Ошибочный Директор Иванович' not in combined


def test_nonstandard_scope_is_visible_in_smk_orders_but_not_added_to_customer_report():
    data = sample_data()
    scope = 'Производство металлоконструкций и разработка проектной документации'
    data['company']['scope'] = scope
    data['objects'] = [
        {'name': 'Цех металлоконструкций «Север»', 'year': '2026', 'customer': 'ООО «Заказчик 1»'},
        {'name': 'Проект склада № 7', 'year': '2026', 'customer': 'ОДО «Заказчик 2»'},
    ]
    itr = [x for x in data['staff'] if not x['is_worker']]
    workers = [x for x in data['staff'] if x['is_worker']]
    result = generate_iso_suot_package_v2(
        data['company'], itr, generator.calculate_dates('20.08.2026'), generator.select_responsible(itr),
        product='iso', workers=workers, suppliers=data['suppliers'], objects=data['objects']
    )

    order = next(d for d in result['docs'] if 'Приказ 3-СМК' in d['name'])
    report = next(d for d in result['docs'] if 'оценке удовлетворенности заказчиков' in d['name'].lower())
    order_text = xml_text(order['bytes'])
    report_text = xml_text(report['bytes'])
    assert scope in order_text
    assert 'В ОБЛАСТИ:' in order_text
    assert scope not in report_text
    assert 'В ОБЛАСТИ:' not in report_text
    assert 'Цех металлоконструкций «Север»' in report_text
    assert 'Проект склада № 7' in report_text
    assert 'ООО «Заказчик 1»' in report_text
    assert 'ОДО «Заказчик 2»' in report_text
    assert 'Устройство системы адресной пожарной сигнализации' not in report_text


def test_iso_scope_from_current_card_never_turns_into_a_placeholder():
    data = sample_data()
    scope = 'Производство строительно-монтажных работ'
    data['company']['scope'] = 'Устаревшая область'
    data['certification']['scope'] = scope

    result = generator.generate_package(data, 'dummy', 'iso_suot')

    assert result['docs']
    order = next(doc for doc in result['docs'] if 'Приказ 3-СМК' in doc['name'])
    order_document = Document(io.BytesIO(order['bytes']))
    order_text = '\n'.join(
        [paragraph.text for paragraph in order_document.paragraphs]
        + [cell.text for table in order_document.tables for row in table.rows for cell in row.cells]
    )
    assert f'В ОБЛАСТИ: {scope}' in order_text
    assert 'В ОБЛАСТИ: ТРЕБУЕТ УТОЧНЕНИЯ' not in order_text


def test_iso_uses_explicit_document_date_and_preserves_approved_forms():
    data = sample_data()
    data['dates'] = {'document_date': '12.01.2026'}
    data['suppliers'] = [
        {'name': 'ООО «Первый поставщик»', 'type': 'Материалы'},
        {'name': 'ООО «Второй поставщик»', 'type': 'Материалы'},
    ]
    data['objects'] = [
        {'name': 'Объект 1', 'year': '2025'},
        {'name': 'Объект 2', 'year': '2025'},
    ]
    result = generator.generate_package(data, 'dummy', 'iso')
    docs = {doc['name']: doc['bytes'] for doc in result['docs']}
    assert not any('повышения квалификации' in name.lower() for name in docs)

    def text(fragment):
        doc = Document(io.BytesIO(fragment))
        return '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(
            '\t'.join(c.text for c in row.cells) for table in doc.tables for row in table.rows
        )

    policy = text(next(value for name, value in docs.items() if 'Политика в области качества' in name))
    familiarization = text(next(value for name, value in docs.items() if 'Лист ознакомления с целями' in name))
    audits = text(next(value for name, value in docs.items() if 'Программа внутренних аудитов СМК' in name))
    supplier = text(next(value for name, value in docs.items() if 'Карточка оценки поставщика 1' in name))
    report = text(next(value for name, value in docs.items() if 'оценке удовлетворенности заказчиков' in name.lower()))
    goals = text(next(value for name, value in docs.items() if '2.1 Цели' in name))
    opportunities = text(next(value for name, value in docs.items() if 'план по возможностям' in name))

    assert '12.01.2026' in policy and 'Варта' not in policy
    assert 'УТВЕРЖДАЮ' not in familiarization
    assert 'Должность' not in familiarization
    assert familiarization.count('12.01.2026') == 3
    assert 'Журнал регистрации внутренних аудитов' in audits
    assert 'на 2026 год' in audits
    assert 'ООО «Первый поставщик»' in supplier
    assert 'ООО «Второй поставщик»' in supplier
    assert 'за 2025 год' in report
    assert 'В ОБЛАСТИ:' not in report
    assert 'Производитель работ' not in goals
    assert 'Производитель работ' not in opportunities

    process = text(next(value for name, value in docs.items() if 'Отчет по процессу строительства' in name))
    assert '0/6' not in process
    assert '0/2' in process


def _visible_doc_text(doc_bytes):
    document = Document(io.BytesIO(doc_bytes))
    return '\n'.join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )


def test_iso_audit_artifacts_use_current_year_and_never_self_audit():
    data = sample_data()
    company = dict(data['company']); company['scope'] = data['certification']['scope']
    itr = [person for person in data['staff'] if not person['is_worker']]
    workers = [person for person in data['staff'] if person['is_worker']]
    result = generate_iso_suot_package_v2(
        company, itr, generator.calculate_dates('20.08.2026'), generator.select_responsible(itr),
        product='iso', workers=workers, suppliers=data['suppliers'], objects=data['objects'],
    )
    docs = {doc['name']: doc['bytes'] for doc in result['docs']}
    audit = next(value for name, value in docs.items() if 'Программа внутренних аудитов СМК' in name)
    audit_document = Document(io.BytesIO(audit))
    assert 'на 2026 год' in _visible_doc_text(audit)
    assert 'на 2025 год' not in _visible_doc_text(audit)
    for row in audit_document.tables[1].rows[3:]:
        audited = row.cells[1].text
        auditor = row.cells[7].text
        assert audited and auditor and audited != auditor
    assert sum('Контрольный лист внутреннего аудита' in name for name in docs) == len(itr)
    assert sum('Отчет по внутреннему аудиту' in name for name in docs) == len(itr)


def test_iso_coordination_protocol_uses_current_staff_without_duplicate_roles():
    company = {
        'name': 'Мавис Групп', 'form': 'ООО', 'city': 'Минск',
        'director_fio': 'Клепицкая Валентина Алексеевна', 'director_position': 'Директор',
        'scope': 'Производство строительно-монтажных работ',
    }
    itr = [
        {'fio': 'Клепицкая Валентина Алексеевна', 'position': 'Директор', 'ot_certificate': True},
        {'fio': 'Падуто Александр Сергеевич', 'position': 'Начальник участка', 'ot_certificate': True},
        {'fio': 'Кульбакин Александр Валерьевич', 'position': 'Бухгалтер'},
    ]
    result = generate_iso_suot_package_v2(
        company, itr, generator.calculate_dates('20.08.2026'), generator.select_responsible(itr),
        product='iso', workers=[], suppliers=[], objects=[],
    )
    protocol = next(doc['bytes'] for doc in result['docs'] if 'протокол кс' in doc['name'].lower())
    text = _visible_doc_text(protocol)
    assert 'Главный специалист' not in text
    assert text.count('Клепицкая Валентина Алексеевна') == 2  # председатель и подпись
    assert 'Начальник участка — Падуто Александр Сергеевич' in text
    assert 'Кульбакин Александр Валерьевич — Бухгалтер' not in text


def test_iso_current_document_versions_replace_stale_template_references():
    data = sample_data()
    data['dates'] = {'document_date': '12.01.2026'}
    result = generator.generate_package(data, 'dummy', 'iso')
    texts = {name: _visible_doc_text(doc['bytes']) for name, doc in ((d['name'], d) for d in result['docs'])}
    manual = next(text for name, text in texts.items() if 'РК СМК' in name)
    info = next(text for name, text in texts.items() if 'Информационный лист БСЦ' in name)
    management_report = next(text for name, text in texts.items() if 'сводн' in name.lower())
    assert 'СТП СМК 02-2025' not in manual
    assert 'СТП СМК 03-2025' not in manual
    assert 'СТП СМК 04-2025' not in manual
    assert not re.search(r'СТП СМК 03-202(?!\d)', info)
    assert 'СТП СМК 04-2025' not in info
    assert 'СТП СМК 03-2026' in info
    assert 'СТП СМК 04-2026' in info
    assert 'запланировано проведение 8 внутренних аудитов' not in management_report
    assert 'Внутренние аудиты проводились директором и производителем работ' not in management_report
    assert 'по 3 направлениям' in management_report


def test_suot_keeps_ppe_form_and_routes_documents_to_working_folders():
    data = sample_data()
    result = generator.generate_package(data, 'dummy', 'suot')
    paths = {doc['name']: doc['archive_path'] for doc in result['docs']}
    journal_path = next(path for name, path in paths.items() if 'Журнал  учета выдачи  инструкции' in name)
    assert journal_path.startswith('СУОТ/4 Журналы/')
    assert any('Карточка учета выдачи СИЗ (бланк)' in name for name in paths)
    assert any(path.startswith('СУОТ/1 Приказы/') and 'Протокол внутреннего обучения' in name
               for name, path in paths.items())
    assert any(path.startswith('СУОТ/1 Приказы/') and 'Программа внутреннего обучения' in name
               for name, path in paths.items())


def test_package_uses_position_not_model_flag_to_keep_workers_out_of_itr_documents():
    data = sample_data()
    data['staff'] = [
        {'fio': 'Глушинский Олег Иванович', 'position': 'Директор', 'is_worker': True},
        {'fio': 'Лукашик Сергей Васильевич', 'position': 'Заместитель директора-главный инженер', 'is_worker': True},
        {'fio': 'Никифоренко Дмитрий Леонидович', 'position': 'Заместитель директора', 'is_worker': False},
        {'fio': 'Брикин Денис Александрович', 'position': 'Подсобный рабочий', 'role': 'itr', 'is_worker': False},
    ]
    data['workers'] = ['Подсобный рабочий']
    data['dates'] = {'development_date': '27.03.2026'}

    result = generator.generate_package(data, 'dummy', 'iso')
    docs = {doc['name']: doc['bytes'] for doc in result['docs']}
    protocol = next(value for name, value in docs.items() if 'Протокол внутреннего обучения' in name)
    text = '\n'.join(
        [paragraph.text for paragraph in Document(io.BytesIO(protocol)).paragraphs]
        + [cell.text for table in Document(io.BytesIO(protocol)).tables for row in table.rows for cell in row.cells]
    )

    assert 'Глушинский Олег Иванович' in text
    assert 'Лукашик Сергей Васильевич' in text
    assert 'Никифоренко Дмитрий Леонидович' in text
    assert 'Брикин Денис Александрович' not in text


def test_iso_audit_journal_date_column_fits_full_date():
    data = sample_data()
    company = dict(data['company'])
    company['scope'] = data['certification']['scope']
    itr = [person for person in data['staff'] if not person['is_worker']]
    workers = [person for person in data['staff'] if person['is_worker']]

    result = generate_iso_suot_package_v2(
        company, itr, generator.calculate_dates('20.08.2026'), generator.select_responsible(itr),
        product='iso', workers=workers, suppliers=data['suppliers'], objects=data['objects'],
    )
    audit = next(doc for doc in result['docs'] if 'Программа внутренних аудитов СМК' in doc['name'])
    document = Document(io.BytesIO(audit['bytes']))

    # 2.1 cm leaves enough room for DD.MM.YYYY plus Word cell margins.
    assert document.tables[1].columns[0].width.twips >= 1190


def test_iso_manual_uses_single_supplied_legal_address():
    data = sample_data()
    data['certification']['scope'] = 'Производство строительно-монтажных работ'

    result = generator.generate_package(data, 'dummy', 'iso')
    manual = next(doc for doc in result['docs'] if 'РК СМК' in doc['name'])
    document = Document(io.BytesIO(manual['bytes']))
    visible_text = '\n'.join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )

    legal_address_line = next(
        line for line in visible_text.splitlines() if line.startswith('Юридический адрес:')
    )
    assert legal_address_line == f"Юридический адрес: {data['company']['address']}"


def test_company_att_itr_is_enriched_from_staff():
    att = {'itr': [{'fio':'Иванов Иван Иванович','position':'Главный инженер'}]}
    staff = [{
        'fio':'Иванов Иван Иванович','position':'Главный инженер',
        'education_full_text':'Высшее, БНТУ, диплом А №123',
        'trudovye_numbers':['ПК №123','Вкладыш №456'],
        'employment_periods':[{'start':'01.01.2010','end':'01.01.2020','position':'Прораб'}],
    }]
    merged = generator._merge_company_att_itr_from_staff(att, staff)
    p = merged['itr'][0]
    assert 'БНТУ' in p['education_full_text']
    assert p['trudovye_numbers'] == ['ПК №123','Вкладыш №456']
    assert p['employment_periods']


def test_iso_fast_archive_skips_ai_reconciliation(tmp_path):
    # text-only ZIP; if fast mode accidentally makes a reconciliation AI call, fail.
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('Список сотрудников.txt', 'Иванов Иван Иванович | Директор\nПетров Петр Петрович | Штукатур')
        z.writestr('Перечень объектов.txt', 'Объект: ремонт офиса')
    old_company = server._extract_company_details
    old_person = server._reconcile_person_summary
    try:
        server._extract_company_details = lambda *a, **k: (_ for _ in ()).throw(AssertionError('AI company reconciliation called'))
        server._reconcile_person_summary = lambda *a, **k: (_ for _ in ()).throw(AssertionError('AI person reconciliation called'))
        result = server.extract_archive_with_vision(zbuf.getvalue(), 'исо.zip', 'dummy', product='iso_suot')
    finally:
        server._extract_company_details = old_company
        server._reconcile_person_summary = old_person
    assert 'Иванов' in result['text']
    assert 'ремонт офиса' in result['analysis_text']


def test_no_legacy_iso_ai_fallback(monkeypatch):
    import generator_iso_suot_templates as mod
    data = sample_data()
    called = {'legacy': False}
    monkeypatch.setattr(mod, 'generate_iso_suot_package_v2', lambda *a, **k: (_ for _ in ()).throw(RuntimeError('boom')))
    monkeypatch.setattr(generator, '_gen_iso', lambda *a, **k: called.__setitem__('legacy', True))
    result = generator.generate_package(data, 'dummy', 'iso')
    assert result.get('error')
    assert called['legacy'] is False
