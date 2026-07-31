"""Regression tests for product-neutral archive extraction v10."""
from __future__ import annotations

import importlib.util
import io
import sys
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

spec = importlib.util.spec_from_file_location("server_v10", ROOT / "server.py")
server = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(server)


def make_docx(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_xlsx(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_spk_archive() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        root = "СПК тест/"
        z.writestr(root + "Реквизиты.docx", make_docx(
            "ООО ТестСПК, УНП 123456789, юридический адрес г. Минск, БИК TESTBY2X"
        ))
        z.writestr(root + "Виды работ.docx", make_docx(
            "Виды работ: монтаж внутренних систем водоснабжения и канализации"
        ))
        z.writestr(root + "Оборудование/Перечень оборудования.xlsx", make_xlsx([
            ["Наименование", "Количество", "Основание"],
            ["Аппарат сварочный", 2, "собственность"],
            ["Кран автомобильный", 1, "договор аренды №15"],
        ]))
        z.writestr(root + "Средства измерений/Реестр СИ.xlsx", make_xlsx([
            ["Средство измерений", "Заводской номер", "Поверка до"],
            ["Манометр", "M-77", "15.05.2027"],
        ]))
        z.writestr(root + "Договоры аренды/Кран.docx", make_docx(
            "Договор аренды крана автомобильного №15 от 01.02.2026"
        ))
        z.writestr(root + "Иванов Иван Иванович/Диплом.docx", make_docx(
            "Иванов Иван Иванович. Диплом АБ 123456, инженер-строитель"
        ))
        z.writestr(root + "Иванов Иван Иванович/Трудовая книжка.docx", make_docx(
            "Трудовая книжка ТК 777. 01.03.2020 принят производителем работ"
        ))
        z.writestr(root + "Фото объекта.heic", b"not parsed")
    return buf.getvalue()


def fake_ai(prompt, api_key, max_tokens=1500):
    if "КОМПАНИИ-ЗАКАЗЧИКА" in prompt:
        return (
            "Название: ООО ТестСПК\nУНП: 123456789\nЮридический адрес: г. Минск\n"
            "Банковские реквизиты: БИК TESTBY2X\nДиректор: не найдено\n"
            "Телефон/email: не найдено"
        )
    return (
        "1) Иванов Иван Иванович\nПаспорт: не найдено\n"
        "Дипломы: АБ 123456, инженер-строитель\n"
        "Трудовая книжка и вкладыши: ТК 777\n"
        "ПЕРИОДЫ РАБОТЫ:\n- 01.03.2020 — по настоящее время | ООО ТестСПК | производитель работ\n"
        "Аттестаты: не найдено\nНЕУВЕРЕННЫЕ ПОЛЯ: нет"
    )


server._simple_ai_call = fake_ai
result = server.extract_archive_with_vision(build_spk_archive(), "spk-test.zip", "test-key")
text = result["text"]

required = [
    "СОСТАВ АРХИВА",
    "ИСХОДНЫЕ ДОКУМЕНТЫ ИЗ АРХИВА",
    "Виды работ: монтаж внутренних систем водоснабжения",
    "Аппарат сварочный",
    "Кран автомобильный",
    "Манометр",
    "M-77",
    "Договор аренды крана",
    "Папка архива: Оборудование",
    "Папка архива: Средства измерений",
    "Иванов Иван Иванович",
    "Фото объекта.heic",
    "формат пока не читается автоматически",
]
missing = [item for item in required if item not in text]
assert not missing, f"SPK archive data was lost: {missing}\n\n{text}"
assert "ИСХОДНЫЕ ДАННЫЕ ISO/СУОТ" not in text
print("OK: SPK source documents and all archive filenames are preserved")
