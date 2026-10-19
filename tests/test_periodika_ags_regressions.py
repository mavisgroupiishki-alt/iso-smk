import io

from docx import Document

import generator
import generator_iso_suot_templates as templates
import server
import zipfile
import http.client
import http.server
import json
import threading


def test_single_docx_upload_http_extracts_staff(monkeypatch):
    monkeypatch.setattr(server.H, '_require_user', lambda self: {'id': 'test', 'role': 'operator'})
    monkeypatch.setattr(server.H, '_same_origin', lambda self: True)
    monkeypatch.setattr(server, 'auth_route_allowed', lambda *args: True)
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = 'ФИО', 'Должность'
    table.cell(1, 0).text, table.cell(1, 1).text = 'Иванов Иван Иванович', 'Директор'
    buf = io.BytesIO()
    doc.save(buf)
    boundary = 'test-boundary'
    body = (f'--{boundary}\r\nContent-Disposition: form-data; name="product"\r\n\r\niso\r\n'.encode()
            + f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="штатное.docx"\r\n'
              'Content-Type: application/octet-stream\r\n\r\n'.encode()
            + buf.getvalue() + f'\r\n--{boundary}--\r\n'.encode())
    httpd = http.server.ThreadingHTTPServer(('127.0.0.1', 0), server.H)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    try:
        conn = http.client.HTTPConnection('127.0.0.1', httpd.server_port, timeout=10)
        conn.request('POST', '/api/extract-text', body, {'Content-Type': f'multipart/form-data; boundary={boundary}'})
        response = conn.getresponse()
        result = json.loads(response.read())
        assert response.status == 200, result
        assert result['structured_data']['_source_company_unverified']
        assert result['structured_data']['staff'][0]['fio'] == 'Иванов Иван Иванович'
        conn.close()
    finally:
        httpd.shutdown()
        httpd.server_close()
        thread.join(timeout=5)


def test_generation_output_names_do_not_collide_between_tasks():
    first = server._generation_zip_name('Компания', '20261019_120000', 'abc12345')
    second = server._generation_zip_name('Компания', '20261019_120000', 'def12345')
    assert first != second
    assert '/' not in server._generation_zip_name('../Компания', '20261019_120000', 'abc12345')
    assert len(server._generation_zip_name('Компания' * 100, '20261019_120000', 'abc12345').encode()) <= 255


def test_same_name_other_tax_id_does_not_import_objects():
    source = ('--- реквизиты.docx ---\nООО «Альфа»\nУНП 123456789\n\n'
              '--- объекты.docx ---\nООО «Альфа»\nУНП 987654321\nНаименование объекта: Чужой объект')
    result = server._iso_suot_archive_card_data(source)
    assert result.get('_source_company_conflict') or not result.get('objects')


def test_two_tax_ids_in_same_line_cannot_silently_import_director():
    source = ('--- реквизиты.docx ---\nООО «Альфа» УНП 123456789\n\n'
              '--- справка.docx ---\nООО «Альфа» УНП 123456789; второй УНП 987654321; директор: Иванов Иван Иванович')
    result = server._iso_suot_archive_card_data(source)
    assert result.get('_source_company_conflict') or result.get('_source_company_unverified')


def test_archive_with_unidentified_document_requires_confirmation():
    source = ('--- реквизиты.docx ---\nООО «Альфа»\nУНП 123456789\n\n'
              '--- данные.docx ---\nРуководитель: Иванов Иван Иванович\nОтдел: Производство')
    assert server._iso_suot_archive_card_data(source).get('_source_company_unverified')


def test_supplier_tax_id_is_not_treated_as_company_conflict():
    source = ('--- реквизиты.docx ---\nООО «Альфа»\nУНП 123456789\n\n'
              '--- поставщики.docx ---\nООО «Поставщик»\nУНП 987654321')
    result = server._iso_suot_archive_card_data(source)
    assert not result.get('_source_company_conflict')
    assert result.get('_source_company_unverified')


def _package():
    company = {'name': 'Тест', 'form': 'ООО', 'director_fio': 'Иванов Иван Иванович',
               'scope': 'производство строительно-монтажных работ'}
    staff = [{'fio': 'Иванов Иван Иванович', 'position': 'Директор'},
             {'fio': 'Петров Петр Петрович', 'position': 'Начальник участка'}]
    return templates.generate_iso_suot_package_v2(
        company, staff, generator.calculate_dates('14.10.2026'),
        generator.select_responsible(staff), product='iso', periodika=True,
        suppliers=[{'name': 'Краска', 'type': 'краска'}, {'name': 'Кабель', 'type': 'кабель'}],
    )


def test_supplier_comparison_does_not_use_unrelated_company():
    result = _package()
    card = next(d for d in result['docs'] if 'Карточка оценки поставщика 1' in d['name'])
    table = Document(io.BytesIO(card['bytes'])).tables[0]
    assert table.rows[4].cells[1].text == 'ТРЕБУЕТ УТОЧНЕНИЯ'


def test_supplier_card_accepts_explicit_comparison():
    card = templates._supplier_card_doc(
        {'name': 'Тест'}, {'name': 'Краска', 'type': 'краска'},
        {'name': 'Магазин краски'}, {}, 1,
    )
    assert Document(io.BytesIO(card)).tables[0].rows[4].cells[1].text == 'Магазин краски'


def test_audit_journal_does_not_invent_performed_dates():
    result = _package()
    program = next(d for d in result['docs'] if 'Программа внутренних аудитов СМК' in d['name'])
    journal = Document(io.BytesIO(program['bytes'])).tables[1]
    assert all(not row.cells[0].text.strip() for row in journal.rows[3:])


def test_policy_does_not_replace_historical_dates_with_document_date():
    result = _package()
    policy = next(d for d in result['docs'] if 'Политика в области качества' in d['name'])
    doc = Document(io.BytesIO(policy['bytes']))
    assert 'ТРЕБУЕТ УТОЧНЕНИЯ' in '\n'.join(p.text for p in doc.paragraphs)
    assert all(not row.cells[2].text.strip() for row in doc.tables[0].rows[1:])


def test_calibration_schedule_has_no_template_equipment():
    result = _package()
    schedule = next(d for d in result['docs'] if 'график' in d['name'].lower() and 'поверк' in d['name'].lower())
    table = Document(io.BytesIO(schedule['bytes'])).tables[0]
    assert all(not cell.text.strip() for row in table.rows[5:] for cell in row.cells)


def test_periodika_hides_service_files_and_puts_cards_in_their_own_folder():
    result = _package()
    assert not result['extra_files']
    assert all('Служебная информация/' not in d['archive_path'] for d in result['docs'])
    cards = [d for d in result['docs'] if 'Карточка оценки поставщика' in d['name']]
    assert all(d['archive_path'].startswith('Карточки поставщиков/') for d in cards)
    assert any('номенклатура дел' in d['name'].lower() for d in result['docs'])
    assert not any('отчет по процессу строительства' in d['name'].lower() for d in result['docs'])


def test_customer_satisfaction_does_not_invent_scores():
    data = templates.render_generic('smk_doc_14.docx', templates._OLD_COMPANY,
                                    {}, {}, {}, company={}, scope_text='', dates={})
    result = templates._ensure_satisfaction_objects(data, [{'name': 'Объект'}])
    table = next(t for t in Document(io.BytesIO(result)).tables
                 if t.rows and 'Объект / заказчик' in t.rows[0].cells[0].text)
    assert table.rows[-1].cells[2].text != '5'


def test_archive_preserves_historical_customers_from_word_table():
    doc = Document()
    doc.add_paragraph('Утверждаю Директор ООО «Тест»')
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ['Заказчик', 'Количество рекламаций', 'Бальная оценка']):
        cell.text = text
    for name in ('Заказчик А', 'Заказчик Б'):
        for cell, text in zip(table.add_row().cells, [name, '—', '5']):
            cell.text = text
    word = io.BytesIO()
    doc.save(word)
    packed = io.BytesIO()
    with zipfile.ZipFile(packed, 'w') as archive:
        archive.writestr('Отчет по оценке удовлетворенности заказчиков.docx', word.getvalue())
    result = server.extract_archive_with_vision(packed.getvalue(), 'source.zip', '', product='iso', package_mode='periodika')
    assert [row['name'] for row in result['structured_data']['iso_suot']['satisfaction_customers']] == ['Заказчик А', 'Заказчик Б']


def test_archive_does_not_import_customers_from_other_companies():
    def report(owner, customer):
        doc = Document()
        doc.add_paragraph(f'Утверждаю директор ООО «{owner}»')
        table = doc.add_table(rows=1, cols=3)
        for cell, text in zip(table.rows[0].cells, ['Заказчик', 'Рекламации', 'Бальная оценка']):
            cell.text = text
        table.add_row().cells[0].text = customer
        output = io.BytesIO()
        doc.save(output)
        return output.getvalue()
    packed = io.BytesIO()
    with zipfile.ZipFile(packed, 'w') as archive:
        archive.writestr('Отчет удовлетворенности 1.docx', report('Компания А', 'Заказчик А'))
        archive.writestr('Отчет удовлетворенности 2.docx', report('Компания А', 'Заказчик Б'))
        archive.writestr('Отчет удовлетворенности 3.docx', report('Компания Б', 'Чужой заказчик'))
    result = server.extract_archive_with_vision(packed.getvalue(), 'source.zip', '', product='iso', package_mode='periodika')
    assert [row['name'] for row in result['structured_data']['iso_suot']['satisfaction_customers']] == ['Заказчик А', 'Заказчик Б']


def test_approved_report_owner_is_not_confused_with_its_customer():
    source = '''--- Отчет по оценке удовлетворенности заказчиков.docx ---
Утверждаю Директор ООО «Организация А» 01.01.2025
Заказчик ООО «Клиент Б» — 5
'''
    assert server._iso_suot_archive_card_data(source)['company']['name'] == 'Организация А'


def test_two_employers_in_staffing_documents_stop_automatic_transfer():
    text = ('--- штат_а.docx ---\nООО «Компания А»\nИванов Иван Иванович | Директор\n\n'
            '--- штат_б.docx ---\nООО «Компания Б»\nПетров Петр Петрович | Инженер')
    result = server._iso_suot_archive_card_data(text)
    assert result.get('_source_company_conflict') is True
    assert not result.get('staff')


def test_same_name_distinct_unp_in_staffing_stops_transfer():
    text = ('--- штат_а.docx ---\nООО «Тест» УНП 123456789\nИванов Иван Иванович | Директор\n\n'
            '--- штат_б.docx ---\nООО «Тест» УНП 987654321\nПетров Петр Петрович | Инженер')
    assert server._iso_suot_archive_card_data(text).get('_source_company_conflict') is True


def test_employee_list_with_other_employer_stops_transfer():
    text = ('--- штат.docx ---\nООО «Альфа»\nПетров Петр Петрович | Инженер\n\n'
            '--- список.docx ---\nООО «Бета»\nИванов Иван Иванович | Директор')
    assert server._iso_suot_archive_card_data(text).get('_source_company_conflict') is True


def test_staff_owner_must_match_selected_company():
    text = ('--- реквизиты.docx ---\nООО «Альфа»\n\n'
            '--- штатное расписание.docx ---\nООО «Бета»\nИванов Иван Иванович | Директор\n\n'
            '--- карточка.docx ---\nООО «Альфа»')
    assert server._iso_suot_archive_card_data(text).get('_source_company_conflict') is True


def test_director_from_other_company_is_not_assigned_to_selected_company():
    text = ('--- реквизиты.docx ---\nООО «Альфа»\n\n'
            '--- договор.docx ---\nООО «Бета»\nДиректор: Иванов Иван Иванович')
    result = server._iso_suot_archive_card_data(text)
    assert not result.get('company', {}).get('director_fio')


def test_objects_from_other_company_are_not_imported():
    text = ('--- реквизиты.docx ---\nООО «Альфа» УНП 123456789\n\n'
            '--- объекты.docx ---\nООО «Бета»\nНаименование объекта: чужой объект')
    result = server._iso_suot_archive_card_data(text)
    assert not result.get('objects')


def test_scope_from_other_company_is_not_imported():
    text = ('--- реквизиты.docx ---\nООО «Альфа» УНП 123456789\n\n'
            '--- приказ.docx ---\nООО «Бета»\nОбласть СМК: Монтаж чужой системы. 1. Контроль')
    result = server._iso_suot_archive_card_data(text)
    assert not result.get('certification', {}).get('scope')


def test_director_not_guessed_from_document_with_multiple_companies():
    text = ('--- реквизиты.docx ---\nООО «Альфа»\nООО «Бета»\nДиректор: Иванов Иван Иванович')
    result = server._iso_suot_archive_card_data(text)
    assert not result.get('company', {}).get('director_fio')


def test_staffing_document_without_identified_owner_requires_confirmation_even_without_parsed_rows():
    result = server._iso_suot_archive_card_data('--- штатное.docx ---\nТаблица персонала.')
    assert result.get('_source_company_unverified') is True


def test_employee_table_is_detected_without_staff_in_filename():
    text = ('--- информация.docx ---\nФИО | Должность\n'
            'Иванов Иван Иванович | Директор')
    result = server._iso_suot_archive_card_data(text)
    assert result.get('_source_company_unverified') is True
    assert result.get('staff')[0]['fio'] == 'Иванов Иван Иванович'


def test_same_name_selected_company_and_staff_unps_must_match():
    text = ('--- реквизиты.docx ---\nООО «Альфа»\nУНП: 111111111\n\n'
            '--- штат.docx ---\nООО «Альфа»\nУНП: 222222222\nИванов Иван Иванович | Директор')
    assert server._iso_suot_archive_card_data(text).get('_source_company_conflict') is True


def test_supplier_list_is_not_staffing_evidence():
    text = ('--- реквизиты.docx ---\nООО «Альфа»\n\n'
            '--- список поставщиков.docx ---\nООО «Бета»\nООО «Гамма»')
    assert not server._iso_suot_archive_card_data(text).get('_source_company_conflict')


def test_anonymous_staff_source_is_not_automatically_transferred():
    text = ('--- реквизиты.docx ---\nООО «Альфа»\n\n'
            '--- штат.docx ---\nИванов Иван Иванович | Директор')
    result = server._iso_suot_archive_card_data(text)
    assert result.get('_source_company_unverified') is True
    assert result.get('staff')[0]['fio'] == 'Иванов Иван Иванович'


def test_satisfaction_keeps_distinct_historical_evaluations():
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ['Заказчик', 'Рекламации', 'Оценка']):
        cell.text = value
    data = io.BytesIO()
    doc.save(data)
    output = templates._ensure_satisfaction_objects(data.getvalue(), [], [
        {'name': 'Заказчик', 'complaints': 'Нет', 'score': '5'},
        {'name': 'Заказчик', 'complaints': 'Была', 'score': '4'},
    ])
    result = Document(io.BytesIO(output))
    assert len(result.tables[0].rows) == 3


def test_two_employers_in_archive_stop_card_transfer():
    packed = io.BytesIO()
    with zipfile.ZipFile(packed, 'w') as archive:
        for suffix, company, name, position in [('а', 'Компания А', 'Иванов Иван Иванович', 'Директор'),
                                                ('б', 'Компания Б', 'Петров Петр Петрович', 'Инженер')]:
            doc = Document()
            doc.add_paragraph(f'ООО «{company}»')
            table = doc.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = 'ФИО'
            table.rows[0].cells[1].text = 'Должность'
            table.add_row().cells[0].text = name
            table.rows[1].cells[1].text = position
            data = io.BytesIO()
            doc.save(data)
            archive.writestr(f'штат_{suffix}.docx', data.getvalue())
    result = server.extract_archive_with_vision(packed.getvalue(), 'source.zip', '', product='iso', package_mode='periodika')
    assert result['structured_data'].get('_source_company_conflict') is True


def test_conflict_inside_nested_archive_stops_card_transfer():
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, 'w') as archive:
        for suffix, company in [('а', 'Компания А'), ('б', 'Компания Б')]:
            doc = Document()
            doc.add_paragraph(f'ООО «{company}»')
            data = io.BytesIO()
            doc.save(data)
            archive.writestr(f'штат_{suffix}.docx', data.getvalue())
    outer = io.BytesIO()
    with zipfile.ZipFile(outer, 'w') as archive:
        archive.writestr('inner.zip', inner.getvalue())
    result = server.extract_archive_with_vision(outer.getvalue(), 'source.zip', '', product='iso', package_mode='periodika')
    assert result['structured_data'].get('_source_company_conflict') is True


def test_first_document_in_nested_archive_is_available_to_staff_parser():
    inner = io.BytesIO()
    doc = Document()
    doc.add_paragraph('ООО «Тест»')
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = 'ФИО'
    table.rows[0].cells[1].text = 'Должность'
    cells = table.add_row().cells
    cells[0].text = 'Иванов Иван Иванович'
    cells[1].text = 'Директор'
    data = io.BytesIO()
    doc.save(data)
    with zipfile.ZipFile(inner, 'w') as archive:
        archive.writestr('штатное.docx', data.getvalue())
    outer = io.BytesIO()
    with zipfile.ZipFile(outer, 'w') as archive:
        archive.writestr('комплект.zip', inner.getvalue())
    result = server.extract_archive_with_vision(outer.getvalue(), 'source.zip', '', product='iso')
    assert result['structured_data']['staff'][0]['fio'] == 'Иванов Иван Иванович'


def test_nested_source_headers_preserve_first_document_boundary():
    text = ('--- комплект.zip (вложенный архив) ---\n\n'
            '--- штатное.docx ---\nООО «Тест»\nИванов Иван Иванович | Директор\n\n'
            '--- реквизиты.docx ---\nООО «Тест»')
    assert any(path == 'штатное.docx' for path, _ in server._archive_document_blocks(text))


def test_same_name_different_legal_form_is_not_a_match():
    def report(form, name):
        doc = Document()
        doc.add_paragraph(f'Утверждаю Директор {form} «Тест»')
        table = doc.add_table(rows=1, cols=3)
        for cell, text in zip(table.rows[0].cells, ['Заказчик', 'Рекламации', 'Оценка']):
            cell.text = text
        table.add_row().cells[0].text = name
        data = io.BytesIO()
        doc.save(data)
        return data.getvalue()
    packed = io.BytesIO()
    with zipfile.ZipFile(packed, 'w') as archive:
        archive.writestr('Отчет удовлетворенности 1.docx', report('ООО', 'Свой заказчик'))
        archive.writestr('Отчет удовлетворенности 2.docx', report('ООО', 'Ещё свой'))
        archive.writestr('Отчет удовлетворенности 3.docx', report('ЗАО', 'Чужой заказчик'))
    result = server.extract_archive_with_vision(packed.getvalue(), 'source.zip', '', product='iso', package_mode='periodika')
    names = [r['name'] for r in result['structured_data']['iso_suot']['satisfaction_customers']]
    assert names == ['Свой заказчик', 'Ещё свой']


def test_same_name_different_unp_blocks_archive_transfer():
    def report(unp):
        doc = Document()
        doc.add_paragraph(f'Утверждаю Директор ООО «Тест» УНП {unp}')
        table = doc.add_table(rows=1, cols=3)
        for cell, value in zip(table.rows[0].cells, ['Заказчик', 'Рекламации', 'Оценка']):
            cell.text = value
        table.add_row().cells[0].text = 'Заказчик'
        data = io.BytesIO()
        doc.save(data)
        return data.getvalue()
    packed = io.BytesIO()
    with zipfile.ZipFile(packed, 'w') as archive:
        archive.writestr('Отчет удовлетворенности 1.docx', report('123456789'))
        archive.writestr('Отчет удовлетворенности 2.docx', report('987654321'))
    result = server.extract_archive_with_vision(packed.getvalue(), 'source.zip', '', product='iso', package_mode='periodika')
    assert result['structured_data']['_source_company_conflict'] is True
    assert not result['structured_data'].get('iso_suot', {}).get('satisfaction_customers')


def _policy_bytes(approval='05.02.2025', date='07.02.2025'):
    doc = Document()
    doc.add_paragraph('ПОЛИТИКА ООО «Тест»')
    doc.add_paragraph(f'Директор __________ {approval}')
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ['ФИО', 'Подпись', 'Дата']):
        cell.text = text
    for cell, text in zip(table.add_row().cells, ['Иванов Иван Иванович', '', date]):
        cell.text = text
    data = io.BytesIO()
    doc.save(data)
    return data.getvalue()


def test_single_policy_document_extracts_original_dates():
    result = server._iso_suot_single_doc_data(_policy_bytes(), 'Политика качества.docx')
    assert result['iso_suot']['policy_evidence']['approval_date'] == '05.02.2025'
    assert result['iso_suot']['policy_evidence']['familiarizations'][0]['date'] == '07.02.2025'


def test_single_policy_conflicting_familiarization_dates_need_review():
    doc = Document(io.BytesIO(_policy_bytes()))
    for cell, text in zip(doc.tables[0].add_row().cells, ['Иванов Иван Иванович', '', '08.02.2025']):
        cell.text = text
    data = io.BytesIO()
    doc.save(data)
    evidence = server._iso_suot_single_doc_data(data.getvalue(), 'Политика качества.docx')['iso_suot']['policy_evidence']
    assert evidence['familiarizations'][0]['date_conflict']
    assert not evidence['familiarizations'][0]['date']


def test_nested_archive_policy_evidence_is_not_lost():
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, 'w') as archive:
        archive.writestr('Политика качества.docx', _policy_bytes())
    outer = io.BytesIO()
    with zipfile.ZipFile(outer, 'w') as archive:
        archive.writestr('previous.zip', inner.getvalue())
    result = server.extract_archive_with_vision(outer.getvalue(), 'source.zip', '', product='iso', package_mode='periodika')
    assert result['structured_data']['iso_suot']['policy_evidence']['approval_date'] == '05.02.2025'


def test_policy_date_conflicts_are_not_silently_replaced():
    packed = io.BytesIO()
    with zipfile.ZipFile(packed, 'w') as archive:
        archive.writestr('Политика 1.docx', _policy_bytes())
        archive.writestr('Политика 2.docx', _policy_bytes('06.02.2025', '08.02.2025'))
    result = server.extract_archive_with_vision(packed.getvalue(), 'source.zip', '', product='iso', package_mode='periodika')
    evidence = result['structured_data']['iso_suot']['policy_evidence']
    assert evidence['date_conflict'] and not evidence['approval_date']
    assert evidence['familiarizations'][0]['date_conflict'] and not evidence['familiarizations'][0]['date']


def test_word_expansion_is_bounded_for_policy_import():
    data = io.BytesIO()
    with zipfile.ZipFile(data, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', b'0' * (32 * 1024 * 1024 + 1))
    assert not server._iso_suot_safe_docx(data.getvalue())


def test_satisfaction_report_keeps_previous_customers_and_new_object():
    data = templates.render_generic('smk_doc_14.docx', templates._OLD_COMPANY,
                                    {}, {}, {}, company={}, scope_text='', dates={})
    result = templates._ensure_satisfaction_objects(data, [{'name': 'Новый объект'}],
        [{'name': 'Заказчик А'}, {'name': 'Заказчик Б'}])
    table = next(t for t in Document(io.BytesIO(result)).tables
                 if t.rows and 'Объект / заказчик' in t.rows[0].cells[0].text)
    rows = [row.cells[0].text for row in table.rows[1:]]
    assert len(rows) == 3
    assert all(any(name in row for row in rows) for name in ('Заказчик А', 'Заказчик Б', 'Новый объект'))


def test_periodika_policy_uses_confirmed_old_dates_and_new_hire_date():
    dates = generator.calculate_dates('01.10.2026')
    data = templates.render_generic('smk_doc_1.docx', templates._OLD_COMPANY,
                                    {}, {}, templates._date_map(dates), company={}, scope_text='', dates=dates)
    staff = [{'fio': 'Иванов Иван Иванович', 'hire_date': '01.02.2020'},
             {'fio': 'Петров Петр Петрович', 'hire_date': '12.03.2026', 'is_new_employee': True},
             {'fio': 'Сидоров Илья Иванович', 'hire_date': '01.01.2020'}]
    evidence = {'approval_date': '05.02.2025', 'familiarizations': [
        {'fio': 'Иванов Иван Иванович', 'date': '07.02.2025'}]}
    result = templates._iso_policy_doc(data, staff, {'document': '01.10.2026'}, True, evidence)
    doc = Document(io.BytesIO(result))
    assert '05.02.2025' in '\n'.join(p.text for p in doc.paragraphs)
    rows = {r.cells[0].text: r.cells[2].text for r in doc.tables[0].rows[1:]}
    assert rows == {'Иванов Иван Иванович': '07.02.2025',
                    'Петров Петр Петрович': '12.03.2026',
                    'Сидоров Илья Иванович': ''}


def test_policy_retains_familiarization_even_when_approval_date_missing():
    dates = generator.calculate_dates('01.10.2026')
    data = templates.render_generic('smk_doc_1.docx', templates._OLD_COMPANY,
                                    {}, {}, templates._date_map(dates), company={}, scope_text='', dates=dates)
    staff = [{'fio': 'Иванов Иван Иванович'}]
    evidence = {'approval_date': '', 'familiarizations': [{'fio': staff[0]['fio'], 'date': '07.02.2025'}]}
    doc = Document(io.BytesIO(templates._iso_policy_doc(data, staff, dates, True, evidence)))
    assert doc.tables[0].rows[1].cells[2].text == '07.02.2025'
    assert 'ТРЕБУЕТ УТОЧНЕНИЯ' in '\n'.join(p.text for p in doc.paragraphs)


def test_policy_retains_former_employee_familiarization():
    dates = generator.calculate_dates('01.10.2026')
    data = templates.render_generic('smk_doc_1.docx', templates._OLD_COMPANY,
                                    {}, {}, templates._date_map(dates), company={}, scope_text='', dates=dates)
    evidence = {'approval_date': '05.02.2025', 'familiarizations': [{'fio': 'Иванов Иван Иванович', 'date': '07.02.2025'}]}
    doc = Document(io.BytesIO(templates._iso_policy_doc(data, [], dates, True, evidence)))
    assert any(r.cells[0].text == 'Иванов Иван Иванович' and r.cells[2].text == '07.02.2025' for r in doc.tables[0].rows[1:])


def test_satisfaction_preserves_confirmed_scores():
    data = templates.render_generic('smk_doc_14.docx', templates._OLD_COMPANY,
                                    {}, {}, {}, company={}, scope_text='', dates={})
    result = templates._ensure_satisfaction_objects(data, [{'name': 'Новый объект'}],
                [{'name': 'Заказчик А', 'complaints': 'Нет', 'score': '4'}])
    table = next(t for t in Document(io.BytesIO(result)).tables if t.rows and 'Объект / заказчик' in t.rows[0].cells[0].text)
    assert [c.text for c in table.rows[1].cells][:3] == ['Заказчик А', 'Нет', '4']
    assert table.rows[2].cells[2].text == ''


def test_satisfaction_source_reads_confirmed_scores():
    doc = Document()
    doc.add_paragraph('Утверждаю директор ООО «Тест»')
    table = doc.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ['Заказчик', 'Рекламации', 'Оценка']):
        cell.text = value
    for cell, value in zip(table.add_row().cells, ['Заказчик А', 'Нет', '4']):
        cell.text = value
    output = io.BytesIO()
    doc.save(output)
    result = server._iso_suot_satisfaction_source(output.getvalue(), 'Отчет удовлетворенности.docx')
    assert result['customers'][0] == {'name': 'Заказчик А', 'complaints': 'Нет', 'score': '4'}


def test_policy_keeps_all_named_workers_of_the_same_profession():
    workers = [{'fio': 'Работников Иван Иванович', 'position': 'Штукатур', 'hire_date': '12.03.2026', 'is_new_employee': True},
               {'fio': 'Работников Петр Петрович', 'position': 'Штукатур', 'hire_date': '13.03.2026', 'is_new_employee': True}]
    company = {'name': 'Тест', 'form': 'ООО', 'director_fio': 'Иванов Иван Иванович',
               'scope': 'производство строительно-монтажных работ'}
    staff = [{'fio': 'Иванов Иван Иванович', 'position': 'Директор'}]
    result = templates.generate_iso_suot_package_v2(
        company, staff, generator.calculate_dates('01.10.2026'), generator.select_responsible(staff),
        product='iso', periodika=True, workers=workers,
        iso_suot={'policy_evidence': {'approval_date': '05.02.2025', 'familiarizations': []}})
    doc = Document(io.BytesIO(next(d['bytes'] for d in result['docs'] if 'Политика в области качества' in d['name'])))
    rows = {r.cells[0].text: r.cells[2].text for r in doc.tables[0].rows[1:]}
    assert rows['Работников Иван Иванович'] == '12.03.2026'
    assert rows['Работников Петр Петрович'] == '13.03.2026'
